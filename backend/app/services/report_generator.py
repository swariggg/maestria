import io
from datetime import datetime
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from sqlalchemy.orm import Session
from app.models.entities import (
    Inscripcion, Usuario, Programa, Convocatoria, CursoPreparatorio,
    ActividadCurso, AsignacionCurso, EntregaEvidencia, EvaluacionEvidencia,
    SeguimientoPostulante, AsignacionTutor
)

def generate_individual_pdf_report(db: Session, inscripcion_id: int) -> io.BytesIO:
    inscripcion = db.query(Inscripcion).filter(Inscripcion.id == inscripcion_id).first()
    if not inscripcion:
        raise ValueError("Inscripción no encontrada.")

    postulante = inscripcion.postulante
    convocatoria = inscripcion.convocatoria
    programa = convocatoria.programa

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=40, leftMargin=40, topMargin=40, bottomMargin=40
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=18,
        leading=22,
        textColor=colors.HexColor('#1E3A8A'),
        alignment=1 # Centered
    )
    subtitle_style = ParagraphStyle(
        'DocSubtitle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=11,
        leading=14,
        textColor=colors.HexColor('#4B5563'),
        alignment=1
    )
    section_title = ParagraphStyle(
        'SectionTitle',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=13,
        leading=16,
        textColor=colors.HexColor('#252533'),
        spaceBefore=12,
        spaceAfter=6
    )
    body_style = ParagraphStyle(
        'Body',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        leading=13,
        textColor=colors.HexColor('#252533')
    )

    story = []

    # Encabezado Institucional
    story.append(Paragraph("UNIVERSIDAD INTERNACIONAL DEL ECUADOR (UIDE)", title_style))
    story.append(Paragraph("Dirección General de Posgrados - APP Maestría", subtitle_style))
    story.append(Paragraph("INFORME INDIVIDUAL DE SEGUIMIENTO ACADÉMICO PREPARATORIO", subtitle_style))
    story.append(Spacer(1, 10))
    story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor('#BFD5FF'), spaceAfter=15))

    # Datos del Postulante y Programa
    tutor_asig = db.query(AsignacionTutor).filter(
        AsignacionTutor.inscripcion_id == inscripcion.id,
        AsignacionTutor.estado == 'ACTIVO'
    ).first()
    tutor_nombre = tutor_asig.tutor.nombre_completo if tutor_asig else "No asignado"

    info_data = [
        [Paragraph("<b>Postulante:</b>", body_style), Paragraph(postulante.nombre_completo, body_style),
         Paragraph("<b>Cédula / Identificación:</b>", body_style), Paragraph(postulante.cedula, body_style)],
        [Paragraph("<b>Correo:</b>", body_style), Paragraph(postulante.email, body_style),
         Paragraph("<b>Teléfono:</b>", body_style), Paragraph(postulante.telefono or "N/A", body_style)],
        [Paragraph("<b>Programa:</b>", body_style), Paragraph(f"{programa.codigo} - {programa.nombre}", body_style),
         Paragraph("<b>Convocatoria:</b>", body_style), Paragraph(convocatoria.nombre, body_style)],
        [Paragraph("<b>Tutor Asignado:</b>", body_style), Paragraph(tutor_nombre, body_style),
         Paragraph("<b>Fecha Inscripción:</b>", body_style), Paragraph(inscripcion.fecha_inscripcion.strftime("%Y-%m-%d"), body_style)],
        [Paragraph("<b>Estado Proceso:</b>", body_style), Paragraph(f"<b>{inscripcion.estado_proceso}</b>", body_style),
         Paragraph("<b>Nivel de Riesgo:</b>", body_style), Paragraph(f"<b>{inscripcion.nivel_riesgo}</b>", body_style)],
    ]

    info_table = Table(info_data, colWidths=[100, 180, 120, 130])
    info_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#F8FAFC')),
        ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
    ]))
    story.append(info_table)
    story.append(Spacer(1, 15))

    # Cursos y Evaluaciones de Actividades
    story.append(Paragraph("1. Actividades y Evaluaciones del Curso Preparatorio", section_title))
    
    act_headers = [
        Paragraph("<b>Actividad / Proyecto</b>", body_style),
        Paragraph("<b>Ponderación</b>", body_style),
        Paragraph("<b>Estado Entrega</b>", body_style),
        Paragraph("<b>Calificación</b>", body_style),
        Paragraph("<b>Resultado</b>", body_style),
    ]
    act_rows = [act_headers]

    asig_curso = db.query(AsignacionCurso).filter(AsignacionCurso.inscripcion_id == inscripcion.id).first()
    if asig_curso:
        actividades = db.query(ActividadCurso).filter(ActividadCurso.curso_id == asig_curso.curso_id).all()
        for act in actividades:
            entrega = db.query(EntregaEvidencia).filter(
                EntregaEvidencia.asignacion_curso_id == asig_curso.id,
                EntregaEvidencia.actividad_id == act.id
            ).first()

            estado_ent = entrega.estado_entrega if entrega else "PENDIENTE"
            calif_str = "S/N"
            res_str = "Pendiente"

            if entrega and entrega.versiones:
                ultima_version = entrega.versiones[-1]
                if ultima_version.evaluacion:
                    calif_str = f"{ultima_version.evaluacion.puntaje_obtenido:.2f}"
                    res_str = ultima_version.evaluacion.resultado

            act_rows.append([
                Paragraph(act.titulo, body_style),
                Paragraph(f"{act.ponderacion}%", body_style),
                Paragraph(estado_ent, body_style),
                Paragraph(calif_str, body_style),
                Paragraph(res_str, body_style),
            ])
    
    act_table = Table(act_rows, colWidths=[200, 75, 95, 80, 80])
    act_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#DCE7FF')),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ('ALIGN', (1, 0), (-1, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
    ]))
    story.append(act_table)
    story.append(Spacer(1, 15))

    # Historial Cronológico de Seguimientos
    story.append(Paragraph("2. Historial de Seguimientos Registrados", section_title))
    seguimientos = db.query(SeguimientoPostulante).filter(
        SeguimientoPostulante.inscripcion_id == inscripcion.id
    ).order_by(SeguimientoPostulante.fecha_contacto.asc()).all()

    if seguimientos:
        seg_headers = [
            Paragraph("<b>Fecha</b>", body_style),
            Paragraph("<b>Canal</b>", body_style),
            Paragraph("<b>Motivo / Acción</b>", body_style),
            Paragraph("<b>Respuesta Postulante</b>", body_style),
            Paragraph("<b>Estado</b>", body_style),
        ]
        seg_rows = [seg_headers]
        for s in seguimientos:
            resp_texto = s.respuesta_postulante or ("Respondido" if s.hubo_respuesta else "SIN RESPUESTA")
            seg_rows.append([
                Paragraph(s.fecha_contacto.strftime("%Y-%m-%d"), body_style),
                Paragraph(s.tipo_contacto, body_style),
                Paragraph(f"{s.motivo}<br/><font color='#6B7280'>{s.accion_realizada}</font>", body_style),
                Paragraph(resp_texto, body_style),
                Paragraph(s.estado_caso, body_style),
            ])
        
        seg_table = Table(seg_rows, colWidths=[70, 80, 180, 130, 70])
        seg_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#F8C7DE')),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('TOPPADDING', (0, 0), (-1, -1), 5),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
        ]))
        story.append(seg_table)
    else:
        story.append(Paragraph("<i>No se registran contactos de seguimiento a la fecha.</i>", body_style))

    story.append(Spacer(1, 25))
    story.append(Paragraph(f"Informe generado automáticamente el {datetime.now().strftime('%Y-%m-%d %H:%M:%S')} - Sistema APP Maestría UIDE.", ParagraphStyle('Footer', parent=styles['Italic'], fontSize=8, textColor=colors.HexColor('#9CA3AF'), alignment=1)))

    doc.build(story)
    buffer.seek(0)
    return buffer

def generate_consolidated_excel_report(db: Session, programa_id: int = None, convocatoria_id: int = None) -> io.BytesIO:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Matriz de Seguimiento"

    # Estilos UIDE
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    title_font = Font(name="Calibri", size=14, bold=True, color="1E3A8A")
    subtitle_font = Font(name="Calibri", size=10, italic=True, color="4B5563")
    border_thin = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    # Título
    ws.merge_cells("A1:I1")
    ws["A1"] = "UNIVERSIDAD INTERNACIONAL DEL ECUADOR - DIRECCIÓN DE POSGRADOS"
    ws["A1"].font = title_font
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells("A2:I2")
    ws["A2"] = f"Matriz Consolidada de Seguimiento Académico - Cursos Preparatorios (Generado: {datetime.now().strftime('%Y-%m-%d %H:%M')})"
    ws["A2"].font = subtitle_font
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")

    # Encabezados
    headers = [
        "Postulante", "Cédula", "Programa", "Convocatoria", "Fecha Inscripción",
        "% Avance", "Tutor Asignado", "# Seguimientos", "Estado", "Nivel de Riesgo"
    ]
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border_thin

    # Query
    query = db.query(Inscripcion)
    if convocatoria_id:
        query = query.filter(Inscripcion.convocatoria_id == convocatoria_id)
    elif programa_id:
        query = query.join(Convocatoria).filter(Convocatoria.programa_id == programa_id)
    
    inscripciones = query.all()

    row_num = 5
    for insc in inscripciones:
        post = insc.postulante
        conv = insc.convocatoria
        prog = conv.programa

        # Tutor
        tutor_asig = db.query(AsignacionTutor).filter(
            AsignacionTutor.inscripcion_id == insc.id,
            AsignacionTutor.estado == 'ACTIVO'
        ).first()
        tutor_nombre = tutor_asig.tutor.nombre_completo if tutor_asig else "Sin Tutor"

        # Seguimientos
        num_seg = db.query(SeguimientoPostulante).filter(SeguimientoPostulante.inscripcion_id == insc.id).count()

        # Avance
        asig_c = db.query(AsignacionCurso).filter(AsignacionCurso.inscripcion_id == insc.id).first()
        pct_avance = 0.0
        if asig_c:
            total_acts = db.query(ActividadCurso).filter(ActividadCurso.curso_id == asig_c.curso_id).count()
            if total_acts > 0:
                aprobadas = db.query(EntregaEvidencia).filter(
                    EntregaEvidencia.asignacion_curso_id == asig_c.id,
                    EntregaEvidencia.estado_entrega == 'APROBADO'
                ).count()
                pct_avance = round((aprobadas / total_acts) * 100.0, 1)

        row_data = [
            post.nombre_completo,
            post.cedula,
            prog.codigo,
            conv.codigo,
            insc.fecha_inscripcion.strftime("%Y-%m-%d"),
            f"{pct_avance}%",
            tutor_nombre,
            num_seg,
            insc.estado_proceso,
            insc.nivel_riesgo
        ]

        for col_idx, val in enumerate(row_data, 1):
            cell = ws.cell(row=row_num, column=col_idx, value=val)
            cell.border = border_thin
            if col_idx in [2, 3, 4, 5, 6, 8, 9, 10]:
                cell.alignment = Alignment(horizontal="center")
            
            # Pintar riesgo
            if col_idx == 10:
                if val == "ALTO":
                    cell.fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
                    cell.font = Font(color="991B1B", bold=True)
                elif val == "MEDIO":
                    cell.fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
                    cell.font = Font(color="92400E", bold=True)
                else:
                    cell.fill = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")
                    cell.font = Font(color="065F46", bold=True)

        row_num += 1

    # Autoajuste de columnas
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
